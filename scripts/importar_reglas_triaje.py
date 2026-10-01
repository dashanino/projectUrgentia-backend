import os
import sys
import unicodedata
import openpyxl

# Permite importar "app" cuando ejecutamos el script desde /scripts
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from app import create_app
from app.models.antecedente import Antecedente
from app.models.poblacion import Poblacion
from app.models.red_flag import RedFlag
from app.extensions import db
from app.models.triage_rule import TriageRule
from app.models.triage_rule_antecedente import TriageRuleAntecedente

ARCHIVO_EXCEL = os.path.join(
    BASE_DIR,
    "combinaciones_ESI_normalizado_BD.xlsx"
)

HOJA = "Combinaciones completas"


def normalizar_texto(valor):
    """
    Normalización SOLO para comparar nombres.
    No modifica lo almacenado en la BD.
    """
    if valor is None:
        return ""

    texto = str(valor).strip().lower()

    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(
        caracter
        for caracter in texto
        if unicodedata.category(caracter) != "Mn"
    )

    return texto


def crear_mapa(registros):
    return {
        normalizar_texto(registro.name): registro
        for registro in registros
    }


def separar_antecedentes(valor):
    if valor is None:
        return []

    texto = str(valor).strip()

    if normalizar_texto(texto) == "ninguno":
        return []

    # El Excel usa "+" cuando hay varios antecedentes
    return [
        antecedente.strip()
        for antecedente in texto.split("+")
        if antecedente.strip()
    ]


def validar_excel():
    app = create_app()

    with app.app_context():

        if not os.path.exists(ARCHIVO_EXCEL):
            print("\n❌ No encontré el Excel:")
            print(ARCHIVO_EXCEL)
            return

        workbook = openpyxl.load_workbook(
            ARCHIVO_EXCEL,
            read_only=True,
            data_only=True
        )

        if HOJA not in workbook.sheetnames:
            print(f"\n❌ No existe la hoja '{HOJA}'")
            return

        sheet = workbook[HOJA]

        encabezados = [
            celda.value.strip() if isinstance(celda.value, str) else celda.value
            for celda in next(sheet.iter_rows(min_row=1, max_row=1))
        ]

        columnas_necesarias = [
            "Antecedentes seleccionados",
            "Tipo de población",
            "Bandera roja",
            "Triaje final",
        ]

        faltantes = [
            columna
            for columna in columnas_necesarias
            if columna not in encabezados
        ]

        if faltantes:
            print("\n❌ Faltan columnas en el Excel:")
            for columna in faltantes:
                print(f"   - {columna}")
            return

        indice = {
            nombre: encabezados.index(nombre)
            for nombre in columnas_necesarias
        }

        poblaciones = crear_mapa(Poblacion.query.all())
        banderas = crear_mapa(RedFlag.query.all())
        antecedentes_bd = crear_mapa(Antecedente.query.all())

        errores = []
        total_filas = 0
        total_si = 0
        total_no = 0

        for numero_fila, fila in enumerate(
            sheet.iter_rows(min_row=2, values_only=True),
            start=2
        ):
            antecedentes_excel = fila[
                indice["Antecedentes seleccionados"]
            ]
            poblacion_excel = fila[
                indice["Tipo de población"]
            ]
            bandera_excel = fila[
                indice["Bandera roja"]
            ]
            triaje_excel = fila[
                indice["Triaje final"]
            ]

            # Ignorar filas completamente vacías
            if not any([
                antecedentes_excel,
                poblacion_excel,
                bandera_excel,
                triaje_excel
            ]):
                continue

            total_filas += 1

            poblacion_key = normalizar_texto(poblacion_excel)
            bandera_key = normalizar_texto(bandera_excel)
            triaje_key = normalizar_texto(triaje_excel)

            if poblacion_key not in poblaciones:
                errores.append(
                    f"Fila {numero_fila}: población no encontrada "
                    f"'{poblacion_excel}'"
                )

            if bandera_key not in banderas:
                errores.append(
                    f"Fila {numero_fila}: bandera no encontrada "
                    f"'{bandera_excel}'"
                )

            for antecedente in separar_antecedentes(
                antecedentes_excel
            ):
                antecedente_key = normalizar_texto(antecedente)

                if antecedente_key not in antecedentes_bd:
                    errores.append(
                        f"Fila {numero_fila}: antecedente no encontrado "
                        f"'{antecedente}'"
                    )

            if triaje_key == "si":
                total_si += 1
            elif triaje_key == "no":
                total_no += 1
            else:
                errores.append(
                    f"Fila {numero_fila}: Triaje final inválido "
                    f"'{triaje_excel}'"
                )

        workbook.close()

        print("\n========== VALIDACIÓN ==========")
        print(f"Filas encontradas: {total_filas}")
        print(f"Triaje Sí: {total_si}")
        print(f"Triaje No: {total_no}")

        if errores:
            print(f"\n❌ Se encontraron {len(errores)} errores:\n")

            for error in errores:
                print(error)

            print(
                "\n⚠️ No se insertó nada en la base de datos."
            )

        else:
            print("\n✅ VALIDACIÓN CORRECTA")
            print("Todas las filas coinciden con la base de datos.")
            print("Todavía NO se insertó ninguna regla.")
def importar_reglas():
    app = create_app()

    with app.app_context():

        # Protección: no importar dos veces
        reglas_existentes = TriageRule.query.count()

        if reglas_existentes > 0:
            print(
                f"\n❌ Ya existen {reglas_existentes} reglas en triage_rules."
            )
            print("No se realizó ninguna importación.")
            return

        workbook = openpyxl.load_workbook(
            ARCHIVO_EXCEL,
            read_only=True,
            data_only=True
        )

        sheet = workbook[HOJA]

        encabezados = [
            celda.value.strip() if isinstance(celda.value, str)
            else celda.value
            for celda in next(
                sheet.iter_rows(min_row=1, max_row=1)
            )
        ]

        indice = {
            nombre: encabezados.index(nombre)
            for nombre in [
                "Antecedentes seleccionados",
                "Tipo de población",
                "Bandera roja",
                "Triaje final",
            ]
        }

        poblaciones = crear_mapa(Poblacion.query.all())
        banderas = crear_mapa(RedFlag.query.all())
        antecedentes_bd = crear_mapa(Antecedente.query.all())

        total_reglas = 0
        total_relaciones = 0

        try:

            for fila in sheet.iter_rows(
                min_row=2,
                values_only=True
            ):

                antecedentes_excel = fila[
                    indice["Antecedentes seleccionados"]
                ]

                poblacion_excel = fila[
                    indice["Tipo de población"]
                ]

                bandera_excel = fila[
                    indice["Bandera roja"]
                ]

                triaje_excel = fila[
                    indice["Triaje final"]
                ]

                # Ignorar filas vacías
                if not any([
                    antecedentes_excel,
                    poblacion_excel,
                    bandera_excel,
                    triaje_excel
                ]):
                    continue

                poblacion = poblaciones[
                    normalizar_texto(poblacion_excel)
                ]

                bandera = banderas[
                    normalizar_texto(bandera_excel)
                ]

                alta_prioridad = (
                    normalizar_texto(triaje_excel) == "si"
                )

                # Crear la regla
                regla = TriageRule(
                    id_poblacion=poblacion.id,
                    id_bandera=bandera.id,
                    alta_prioridad=alta_prioridad
                )

                db.session.add(regla)

                # Necesitamos el ID antes de crear
                # las relaciones con antecedentes
                db.session.flush()

                nombres_antecedentes = separar_antecedentes(
                    antecedentes_excel
                )

                for nombre in nombres_antecedentes:

                    antecedente = antecedentes_bd[
                        normalizar_texto(nombre)
                    ]

                    relacion = TriageRuleAntecedente(
                        id_regla=regla.id,
                        id_antecedente=antecedente.id
                    )

                    db.session.add(relacion)

                    total_relaciones += 1

                total_reglas += 1

            # Un solo commit para toda la importación
            db.session.commit()

            print("\n========== IMPORTACIÓN ==========")
            print(f"Reglas insertadas: {total_reglas}")
            print(
                "Relaciones con antecedentes insertadas: "
                f"{total_relaciones}"
            )
            print("\n✅ IMPORTACIÓN COMPLETADA")

        except Exception as error:

            db.session.rollback()

            print("\n❌ ERROR DURANTE LA IMPORTACIÓN")
            print(error)
            print(
                "\n⚠️ Se hizo rollback. "
                "No quedó una importación parcial."
            )

        finally:
            workbook.close()


if __name__ == "__main__":

    validar_excel()

    respuesta = input(
        "\n¿Deseas importar las reglas a la base de datos? (s/n): "
    )

    if respuesta.strip().lower() == "s":
        importar_reglas()
    else:
        print("\nImportación cancelada.")